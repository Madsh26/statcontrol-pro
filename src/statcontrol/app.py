"""Aplicación Dash con estado aislado por pestaña y resultados invalidables."""

import base64
import io
import logging
from pathlib import Path

from dash import Dash, Input, Output, State, ctx, dcc, html, no_update
from dash.exceptions import PreventUpdate
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .data import (
    decode_upload,
    demo_frame,
    frame_store,
    read_project,
    read_table,
    restore_frame,
    suggest_columns,
)
from .figures import control_figure, distribution_figures, empty, pareto_figure
from .reports import csv_bytes, excel_bytes, pdf_bytes, project_bytes, result_frame
from .statistics import analyze
from .ui import layout, metric, table

log = logging.getLogger(__name__)
CONFIG_IDS = [
    "chart",
    "layout",
    "decimal",
    "value",
    "group",
    "size",
    "measures",
    "frequency",
    "baseline",
    "lsl",
    "usl",
    "rules",
]


def notice(message, kind="info"):
    return html.Div(message, className=f"notice {kind}")


def config_values(values):
    config = dict(zip(CONFIG_IDS, values))
    config["rules"] = "extended" in (config["rules"] or [])
    return config


def create_app():
    app = Dash(
        __name__,
        assets_folder=str(Path(__file__).parent / "assets"),
        title="StatControl Pro | Control de procesos",
        update_title="Analizando…",
    )
    app.server.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
    app.layout = layout

    @app.callback(
        Output("layout", "disabled"),
        Output("group", "disabled"),
        Output("size", "disabled"),
        Output("measures", "disabled"),
        Output("frequency", "disabled"),
        Input("chart", "value"),
        Input("layout", "value"),
        Input("group", "value"),
    )
    def relevant_fields(chart, arrangement, group):
        grouped = chart in {"xr", "xs"}
        return (
            not grouped,
            not grouped,
            not grouped or arrangement == "wide" or bool(group),
            not grouped or arrangement != "wide",
            chart != "pareto",
        )

    @app.callback(
        Output("source", "data"),
        Output("load-status", "children"),
        Input("upload", "contents"),
        Input("demo", "n_clicks"),
        Input("manual-load", "n_clicks"),
        Input("project-upload", "contents"),
        State("upload", "filename"),
        State("manual", "value"),
        prevent_initial_call=True,
    )
    def load(contents, demo, manual_click, project_contents, filename, manual):
        try:
            trigger = ctx.triggered_id
            if trigger == "demo":
                raw = demo_frame().to_csv(index=False).encode()
                filename = "Ejemplo · Diámetro de piezas.csv"
            elif trigger == "manual-load":
                if not manual or not manual.strip():
                    raise ValueError("Pega tus datos antes de continuar.")
                raw = manual.encode("utf-8")
                filename = "Datos manuales.tsv" if "\t" in manual else "Datos manuales.csv"
            elif trigger == "project-upload":
                project = read_project(decode_upload(project_contents))
                return {"project": project}, notice(
                    "Proyecto abierto. Revisa la configuración restaurada y pulsa Analizar proceso.",
                    "success",
                )
            else:
                raw = decode_upload(contents)
            return {"raw": base64.b64encode(raw).decode(), "filename": filename}, notice(
                f"Fuente cargada: {filename}. Revisa la lectura antes de analizar.", "success"
            )
        except Exception as exc:
            log.info("Importación rechazada: %s", exc)
            return None, notice("No se pudo cargar: " + str(exc), "error")

    @app.callback(
        Output("dataset", "data"),
        Output("preview", "children"),
        Output("read-notes", "children"),
        Output("data-summary", "children"),
        Output("sheet", "options"),
        Input("source", "data"),
        Input("separator", "value"),
        Input("header", "value"),
        Input("sheet", "value"),
    )
    def preview(source, separator, header, sheet):
        if not source:
            return (
                None,
                html.Div("Carga un archivo o explora el ejemplo.", className="empty-state"),
                [],
                "",
                [],
            )
        try:
            if "project" in source:
                data = source["project"]["data"]
                frame = restore_frame(data)
                notes = [
                    "Proyecto restaurado: la lectura original se conserva. Los controles de separador y hoja no modifican un proyecto JSON."
                ]
                data = frame_store(frame, data.get("name", "Proyecto"), notes)
                data["saved_config"] = source["project"]["config"]
                sheets = []
            else:
                imported = read_table(
                    base64.b64decode(source["raw"]), source["filename"], separator, header, sheet
                )
                frame, notes, sheets = imported.frame, imported.notes, imported.sheets
                data = frame_store(frame, source["filename"], notes)
            return (
                data,
                table(frame),
                [notice(n) for n in notes],
                f"{len(frame):,} filas · {len(frame.columns)} columnas · Mostrando las primeras {min(12, len(frame))} filas",
                [{"label": s, "value": s} for s in sheets],
            )
        except Exception as exc:
            return (
                None,
                html.Div("No hay una lectura válida todavía.", className="empty-state"),
                notice(str(exc), "error"),
                "",
                [],
            )

    @app.callback(
        *[Output(id, "options") for id in ["value", "group", "measures", "frequency"]],
        *[Output(id, "value") for id in CONFIG_IDS],
        Output("run", "disabled"),
        Input("dataset", "data"),
    )
    def configure(data):
        if not data:
            return (
                [[], [], [], []]
                + ["xr", "long", ".", None, None, 5, [], None, None, None, None, ["extended"]]
                + [True]
            )
        frame = restore_frame(data)
        value, group, measures, arrangement, decimal = suggest_columns(frame)
        defaults = dict(
            zip(
                CONFIG_IDS,
                ["xr", arrangement, decimal, value, group, 5, measures, None, None, None, None, True],
            )
        )
        saved = data.get("saved_config", {})
        defaults.update({k: v for k, v in saved.items() if k in CONFIG_IDS})
        defaults["rules"] = ["extended"] if defaults["rules"] else []
        options = [{"label": c, "value": c} for c in frame]
        return [options] * 4 + [defaults[k] for k in CONFIG_IDS] + [False]

    @app.callback(
        Output("result", "data"),
        Output("analysis-status", "children"),
        Output("nav", "value"),
        Input("run", "n_clicks"),
        Input("dataset", "data"),
        *[Input(id, "value") for id in CONFIG_IDS],
        prevent_initial_call=True,
    )
    def calculate(clicks, data, *values):
        if ctx.triggered_id != "run":
            return None, [], no_update
        if not data:
            return None, notice("Primero carga datos válidos.", "error"), no_update
        try:
            result = analyze(restore_frame(data), config_values(values))
            return result, [], "analysis"
        except (ValueError, TypeError, KeyError, OverflowError) as exc:
            return None, notice(str(exc), "error"), no_update
        except Exception:
            log.exception("Error inesperado en el análisis")
            return (
                None,
                notice("No se pudo completar el análisis. Revisa el formato y la configuración.", "error"),
                no_update,
            )

    @app.callback(
        *[Output(id + "-section", "style") for id in ["data", "analysis", "capacity", "reports", "guide"]],
        Input("nav", "value"),
    )
    def navigate(page):
        return [
            {} if page == id else {"display": "none"}
            for id in ["data", "analysis", "capacity", "reports", "guide"]
        ]

    @app.callback(
        Output("metrics", "children"),
        Output("primary-graph", "figure"),
        Output("secondary-graph", "figure"),
        Output("secondary-wrap", "style"),
        Output("signals", "children"),
        Output("results-table", "children"),
        Output("warnings", "children"),
        Output("capacity-content", "children"),
        Output("histogram", "figure"),
        Output("qq", "figure"),
        *[Output(id, "disabled") for id in ["export-pdf", "export-xlsx", "export-csv", "export-project"]],
        Input("result", "data"),
    )
    def display(result):
        if not result:
            return (
                [
                    metric(
                        "ANÁLISIS PENDIENTE", "—", "Carga datos o vuelve a analizar tras modificar parámetros"
                    )
                ],
                empty(),
                empty(),
                {},
                "Aún no hay resultados vigentes.",
                [],
                [],
                "Agrega especificaciones y ejecuta un análisis.",
                empty(),
                empty(),
                True,
                True,
                True,
                True,
            )
        if result["chart"] == "pareto":
            return (
                [
                    metric("OCURRENCIAS", result["total"], "Frecuencia total"),
                    metric("CATEGORÍAS", len(result["labels"]), "Ordenadas por frecuencia"),
                ],
                pareto_figure(result),
                empty(),
                {"display": "none"},
                "Pareto prioriza categorías; no evalúa estabilidad estadística.",
                table(result_frame(result), 20),
                [],
                notice("La capacidad no se aplica a un diagrama de Pareto."),
                empty(),
                empty(),
                False,
                False,
                False,
                False,
            )
        alert_points = len({a["index"] for a in result["signals"]})
        metrics = [
            metric("OBSERVACIONES", result["total"], "Mediciones válidas"),
            metric(
                "PUNTOS / SUBGRUPOS",
                len(result["labels"]),
                f"{result['baseline']} en referencia · n = {result['n']}",
            ),
            metric("PUNTOS CON SEÑAL", alert_points, f"{len(result['signals'])} eventos detectados"),
            metric("SIGMA INTERNA", f"{result['sigma']:.4f}", "Estimada sobre la referencia", True),
        ]
        alerts = [
            html.Div(
                [
                    html.Span(a["rule"], className="signal-tag"),
                    html.Strong(f"{a['chart']} · Punto {a['point']}"),
                    html.Span(a["detail"]),
                ],
                className="signal-row",
            )
            for a in result["signals"][:100]
        ]
        if len(result["signals"]) > 100:
            alerts.append(notice("Se muestran 100 señales. Exporta Excel para consultar todas."))
        cap = result.get("capability")
        capacity_content = notice(
            "Introduce uno o ambos límites de especificación en Datos y vuelve a analizar."
        )
        if cap and not cap.get("unavailable"):
            capacity_content = [
                html.Div(
                    [
                        metric(
                            k.upper(),
                            f"{cap[k]:.3f}" if cap[k] is not None else "—",
                            "Variación interna" if k in {"cp", "cpk"} else "Variación global",
                        )
                        for k in ["cp", "cpk", "pp", "ppk"]
                    ],
                    className="metrics",
                ),
                notice(
                    "Referencia sin señales según las reglas seleccionadas. Revisa los supuestos antes de concluir capacidad."
                    if cap["stable"]
                    else "Referencia con señales: los índices son descriptivos, no una validación de capacidad.",
                    "info" if cap["stable"] else "warning",
                ),
                html.Div(
                    [
                        html.P(
                            f"Fuera de especificación observado: {cap['observed_pct']:.2f}% · PPM del modelo normal: {cap['normal_ppm']:.1f}"
                        ),
                        html.P(
                            f"Shapiro–Wilk p = {cap['normality_p']:.4f}"
                            if cap["normality_p"] is not None
                            else "Shapiro–Wilk no calculado: se requieren de 3 a 5.000 observaciones."
                        ),
                        html.P(
                            "Los histogramas, índices y prueba de normalidad describen exclusivamente los datos de referencia."
                        ),
                    ],
                    className="notice",
                ),
                *[notice(w, "warning") for w in result["warnings"]],
            ]
        elif cap:
            capacity_content = notice(cap["unavailable"], "warning")
        histogram, qq = distribution_figures(result)
        return (
            metrics,
            control_figure(result),
            control_figure(result, True),
            {},
            alerts
            or notice(
                "Sin señales con las reglas seleccionadas. Esto no demuestra por sí solo que el proceso esté estable.",
                "success",
            ),
            table(result_frame(result).round(4), 20),
            [notice(w, "warning") for w in result["warnings"]],
            capacity_content,
            histogram,
            qq,
            False,
            False,
            False,
            False,
        )

    @app.callback(
        Output("download", "data"),
        Output("export-status", "children"),
        *[Input(id, "n_clicks") for id in ["export-pdf", "export-xlsx", "export-csv", "export-project"]],
        State("result", "data"),
        State("dataset", "data"),
        prevent_initial_call=True,
    )
    def export(pdf, xlsx, csv, project, result, data):
        if not result or not data:
            raise PreventUpdate
        try:
            trigger = ctx.triggered_id
            if trigger == "export-pdf":
                payload, filename = pdf_bytes(result, data["name"]), "statcontrol-informe.pdf"
            elif trigger == "export-xlsx":
                payload, filename = excel_bytes(result), "statcontrol-resultados.xlsx"
            elif trigger == "export-csv":
                payload, filename = csv_bytes(result), "statcontrol-resultados.csv"
            else:
                payload, filename = project_bytes(data, result["config"]), "statcontrol-proyecto.json"
            return dcc.send_bytes(payload, filename), notice("Archivo preparado para descargar.", "success")
        except Exception:
            log.exception("Error de exportación")
            return no_update, notice(
                "No se pudo generar el archivo. Los resultados siguen disponibles.", "error"
            )

    @app.callback(
        Output("template-download", "data"),
        Input("template-csv", "n_clicks"),
        Input("template-xlsx", "n_clicks"),
        prevent_initial_call=True,
    )
    def template(csv, xlsx):
        frame = demo_frame()
        if ctx.triggered_id == "template-csv":
            return dcc.send_bytes(frame.to_csv(index=False).encode("utf-8-sig"), "plantilla-larga.csv")
        book = Workbook()
        long = book.active
        long.title = "Largo"
        long.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            long.append(list(row))
        wide = book.create_sheet("Ancho")
        wide.append(["subgrupo", "m1", "m2", "m3", "m4", "m5"])
        for group, rows in frame.groupby("subgrupo", sort=False):
            wide.append([int(group)] + rows.medicion.tolist())
        guide = book.create_sheet("Instrucciones")
        for row in [
            ["STATCONTROL PRO · PLANTILLA"],
            ["Elige Largo o Ancho; no importes esta hoja."],
            ["Sustituye los datos de ejemplo por tus mediciones."],
            ["Largo: una medición por fila; selecciona medicion y subgrupo."],
            ["Ancho: una fila por subgrupo; selecciona m1 a m5 como mediciones."],
            ["Mantén igual tamaño por subgrupo y el orden de medición."],
            ["Sin celdas vacías, títulos adicionales ni separadores de miles."],
            ["Esta muestra es sintética e incluye un desplazamiento para demostrar alertas."],
        ]:
            guide.append(row)
        for sheet in book:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2255B8")
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = 24 if sheet != guide else 100
        buffer = io.BytesIO()
        book.save(buffer)
        return dcc.send_bytes(buffer.getvalue(), "plantilla-statcontrol.xlsx")

    return app


def main():
    create_app().run(host="127.0.0.1", port=8050, debug=False)


if __name__ == "__main__":
    main()
