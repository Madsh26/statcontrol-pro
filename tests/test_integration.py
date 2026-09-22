import io
import json

import pandas as pd
from openpyxl import load_workbook

from statcontrol.app import create_app
from statcontrol.data import demo_frame, frame_store, read_project, restore_frame
from statcontrol.figures import control_figure, distribution_figures, pareto_figure
from statcontrol.reports import csv_bytes, excel_bytes, pdf_bytes, project_bytes
from statcontrol.statistics import analyze


def test_exports_and_reproducibility():
    frame = demo_frame()
    config = dict(
        chart="xr",
        layout="long",
        decimal=".",
        value="medicion",
        group="subgrupo",
        baseline=20,
        lsl=99,
        usl=101,
        rules=True,
    )
    result = analyze(frame, config)
    assert pdf_bytes(result, "example.csv").startswith(b"%PDF")
    book = load_workbook(io.BytesIO(excel_bytes(result)))
    assert book.sheetnames == ["Resultados", "Señales", "Capacidad", "Configuración"]
    assert book["Resultados"].max_row == 26
    assert len(pd.read_csv(io.BytesIO(csv_bytes(result)))) == 25
    project = read_project(project_bytes(frame_store(frame, "example.csv"), config))
    assert analyze(restore_frame(project["data"]), project["config"])["upper"] == result["upper"]
    for fig in [control_figure(result), control_figure(result, True), *distribution_figures(result)]:
        assert len(json.loads(fig.to_json())["data"]) >= 1


def test_pareto_figures_and_formula_safety():
    result = analyze(pd.DataFrame({"category": ["=1+1", "B"]}), {"chart": "pareto", "value": "category"})
    assert len(pareto_figure(result).data) == 2
    book = load_workbook(io.BytesIO(excel_bytes(result)))
    assert book["Resultados"]["A2"].data_type == "s"
    assert book["Resultados"]["A2"].value.startswith("'")
    assert pdf_bytes(result, "pareto.csv").startswith(b"%PDF")


def test_app_routes_and_ids():
    app = create_app()
    client = app.server.test_client()
    assert client.get("/").status_code == 200
    response = client.get("/_dash-layout")
    assert response.status_code == 200
    assert client.get("/_dash-dependencies").status_code == 200
    ids = []
    classes = []

    def walk(obj):
        if isinstance(obj, dict):
            if "props" in obj and "id" in obj["props"]:
                ids.append(obj["props"]["id"])
            if "props" in obj and "className" in obj["props"]:
                classes.append(obj["props"]["className"])
            for value in obj.values():
                walk(value)
        elif isinstance(obj, list):
            for value in obj:
                walk(value)

    walk(response.json)
    assert len(ids) == len(set(ids))
    assert {"upload", "template-csv", "template-xlsx", "run", "guide-section"}.issubset(ids)
    assert "nav-scroll" not in ids
    assert classes.count("capacity-graph") == 2


def test_callbacks_compute_invalidate_and_isolate_sessions():
    app = create_app()
    client = app.server.test_client()
    client.get("/")
    key = next(k for k in app.callback_map if k.startswith("..result.data"))
    meta = app.callback_map[key]
    values = {
        "run": 1,
        "dataset": frame_store(demo_frame(), "demo.csv"),
        "chart": "xr",
        "layout": "long",
        "decimal": ".",
        "value": "medicion",
        "group": "subgrupo",
        "size": 5,
        "measures": ["medicion"],
        "frequency": None,
        "baseline": 20,
        "lsl": 99,
        "usl": 101,
        "rules": ["extended"],
    }

    def request(trigger, browser_client):
        return browser_client.post(
            "/_dash-update-component",
            json={
                "output": key,
                "outputs": [o.to_dict() for o in meta["output"]],
                "inputs": [{**i, "value": values[i["id"]]} for i in meta["inputs"]],
                "state": [],
                "changedPropIds": [trigger],
            },
        )

    computed = request("run.n_clicks", client)
    assert computed.status_code == 200
    assert computed.json["response"]["result"]["data"]["total"] == 125
    assert computed.json["response"]["nav"]["value"] == "analysis"
    invalidated = request("lsl.value", client)
    assert invalidated.json["response"]["result"]["data"] is None
    values["dataset"] = None
    second_client = app.server.test_client()
    isolated = request("run.n_clicks", second_client)
    assert isolated.json["response"]["result"]["data"] is None
